"""Monta o e-mail (HTML + texto) a partir da planilha já recalculada."""
import datetime as dt, html
from openpyxl import load_workbook

CORAL, CORAL_L, INK, INK2, MUTED, BG, RED, GREEN = '#D9644A', '#FBE9E4', '#2A2A33', '#55556A', '#8A8A99', '#F7F6F8', '#B23A3A', '#2E7D5B'

def fmt(v, nf):
    if v is None or v == '': return '–'
    if isinstance(v, str): return v
    if 'R$' in (nf or ''): return 'R$ ' + f'{v:,.0f}'.replace(',', '.')
    if '%' in (nf or ''):
        s = f'{v*100:.0f}%'
        return ('+' + s if v > 0 else s) if '+' in (nf or '') else s
    if isinstance(v, float) and not v.is_integer(): return f'{v:.1f}'.replace('.', ',')
    return f'{v:,.0f}'.replace(',', '.')

def tiles(ws, r):
    out = []
    for c in range(2, 14, 2):
        lab = ws.cell(r, c).value
        if not lab: continue
        v = ws.cell(r + 1, c); sub = ws.cell(r + 2, c).value
        out.append((str(lab).capitalize(), fmt(v.value, v.number_format), sub or '', ws.cell(r, c).fill.fgColor.rgb in ('00FBE9E4', 'FFFBE9E4')))
    return out

def h_tiles(title, ts):
    cells = ''
    for i, (lab, val, sub, acc) in enumerate(ts):
        bg = CORAL_L if acc else BG; col = CORAL if acc else INK
        cells += (f'<td style="width:33%;padding:10px 12px;background:{bg};border:3px solid #fff;vertical-align:top">'
                  f'<div style="font-size:10px;color:{INK2};font-weight:bold;text-transform:uppercase;letter-spacing:.3px">{html.escape(lab)}</div>'
                  f'<div style="font-size:20px;font-weight:bold;color:{col};margin:2px 0">{html.escape(val)}</div>'
                  f'<div style="font-size:11px;color:{MUTED}">{html.escape(str(sub))}</div></td>')
        if i % 3 == 2: cells += '</tr><tr>'
    return (f'<div style="font-size:14px;font-weight:bold;color:{INK};padding:14px 0 6px;border-bottom:2px solid {CORAL};margin-bottom:4px">{html.escape(title)}</div>'
            f'<table style="width:100%;border-collapse:collapse"><tr>{cells}</tr></table>')

def t_tiles(title, ts):
    return f'\n{title.upper()}\n' + '\n'.join(f'- {lab}: {val}' + (f' ({sub})' if sub else '') for lab, val, sub, _ in ts) + '\n'

def h_table(title, head, rows, align=None):
    th = ''.join(f'<td style="padding:5px 8px;font-size:11px;color:{MUTED};text-align:{"left" if i == 0 else "right"}">{html.escape(h)}</td>' for i, h in enumerate(head))
    tr = ''
    for k, r in enumerate(rows):
        bg = BG if k % 2 else '#fff'
        tr += f'<tr style="background:{bg}">' + ''.join(
            f'<td style="padding:6px 8px;font-size:13px;text-align:{"left" if i == 0 else "right"};{"font-weight:bold;" if i == 0 else ""}">{html.escape(str(c))}</td>' for i, c in enumerate(r)) + '</tr>'
    return (f'<div style="font-size:14px;font-weight:bold;color:{INK};padding:14px 0 6px;border-bottom:2px solid {CORAL}">{html.escape(title)}</div>'
            f'<table style="width:100%;border-collapse:collapse"><tr>{th}</tr>{tr}</table>')

def t_table(title, head, rows):
    return f'\n{title.upper()}\n' + '\n'.join('- ' + ' | '.join(str(c) for c in r) for r in rows) + '\n'

def compose(xlsx, stores, modo, hoje, semanal, avisos):
    wb = load_workbook(xlsx, data_only=True)
    ref = stores[0]['ref']
    if modo == 'tarde':
        assunto = f'Encanto · Parcial de hoje {hoje:%d/%m} (14h)'
        intro = f'Parcial do dia {hoje:%d/%m}, com dados da Trinks até {stores[0]["gerado"][11:16]}.'
    else:
        assunto = f'Encanto · Relatório de {hoje:%d/%m}' + (' + fechamento da semana' if semanal else '') + f' (dados até {ref:%d/%m})'
        intro = f'Fechamento até {ref:%d/%m} e o plano para hoje, {hoje:%d/%m}.'
    H = [f'<div style="font-family:Arial,Helvetica,sans-serif;max-width:660px;margin:0 auto;color:{INK}">',
         f'<div style="background:{INK};color:#fff;padding:16px 18px"><div style="font-size:18px;font-weight:bold">Encanto Progressivas</div>'
         f'<div style="font-size:12px;opacity:.85;margin-top:4px">{html.escape(intro)}</div></div>']
    T = [f'Encanto Progressivas. {intro}\n']
    for s in stores:
        W = wb[f'Painel {s["sig"]}']; E = wb[f'Equipe {s["sig"]}']; S = wb[f'Semanas {s["sig"]}']
        H.append(f'<div style="font-size:16px;font-weight:bold;color:{CORAL};padding:16px 0 0">{html.escape(s["loja"])}</div>')
        T.append(f'\n===== {s["loja"].upper()} =====')
        blocks = [(W.cell(9, 2).value, tiles(W, 10)), (W.cell(4, 2).value, tiles(W, 5))] if modo == 'tarde' else \
                 [(W.cell(4, 2).value, tiles(W, 5)), (W.cell(9, 2).value, tiles(W, 10)), (W.cell(14, 2).value, tiles(W, 15))]
        for ttl, ts in blocks:
            H.append(h_tiles(ttl, ts)); T.append(t_tiles(ttl, ts))
        if semanal:
            rows = []
            for r in range(5, 12):
                nm = S.cell(r, 1).value
                if not nm or not str(nm).startswith('Semana'): continue
                if S.cell(r, 12).value == 'futura': continue
                rows.append([f'{nm} ({S.cell(r,2).value:%d/%m}–{S.cell(r,3).value:%d/%m})', fmt(S.cell(r, 6).value, 'R$'), fmt(S.cell(r, 5).value, 'R$'),
                             fmt(S.cell(r, 7).value, '%'), fmt(S.cell(r, 8).value, ''), S.cell(r, 12).value])
            hd = ['Semana', 'Realizado', 'Meta', '%', 'Atend.', 'Situação']
            H.append(h_table('Semana a semana', hd, rows)); T.append(t_table('Semana a semana', hd, rows))
            # colaboradoras na última semana fechada
            wk = None
            for r in range(5, 12):
                if S.cell(r, 12).value in ('✔ bateu', '✖ abaixo'): wk = r - 5
            if wk is not None:
                base = None
                for r in range(1, S.max_row + 1):
                    if S.cell(r, 1).value == 'Colaboradora': base = r
                prs = []
                r = base + 1
                while S.cell(r, 1).value:
                    prs.append([S.cell(r, 1).value, fmt(S.cell(r, 2 + 2 * wk).value, ''), fmt(S.cell(r, 3 + 2 * wk).value, 'R$')]); r += 1
                ttl = f'Equipe na semana {wk+1} ({S.cell(5+wk,2).value:%d/%m}–{S.cell(5+wk,3).value:%d/%m})'
                H.append(h_table(ttl, ['Colaboradora', 'Atendimentos', 'Faturado'], prs)); T.append(t_table(ttl, [], prs))
        rows = []
        for r in range(5, 30):
            nm = E.cell(r, 1).value
            if not nm: break
            if E.cell(r, 11).value in (None, ''): continue
            nf = E.cell(r, 12).number_format
            rows.append([f'{nm} · {str(E.cell(r, 2).value).lower()}', fmt(E.cell(r, 3).value, ''), fmt(E.cell(r, 7).value, 'R$'),
                         f'{E.cell(r, 10).value}: {fmt(E.cell(r, 12).value, nf)}/{fmt(E.cell(r, 11).value, nf)}', fmt(E.cell(r, 13).value, '%')])
        hd = ['Colaboradora', 'Atend.', 'Faturado', 'Meta principal', '%']
        H.append(h_table('Equipe no mês', hd, rows)); T.append(t_table('Equipe no mês', hd, rows))
        ins = [W.cell(21 + k, 9).value for k in range(4) if W.cell(21 + k, 9).value]
        if ins:
            H.append(f'<div style="margin:14px 0;padding:12px 14px;background:{BG};font-size:13px;line-height:1.5"><b>Leitura rápida</b><br>' +
                     '<br>'.join('• ' + html.escape(x) for x in ins) + '</div>')
            T.append('\nLEITURA RÁPIDA\n' + '\n'.join('- ' + x for x in ins))
    if avisos:
        H.append(f'<div style="margin:10px 0;padding:10px 14px;border-left:3px solid {RED};font-size:12px">' + '<br>'.join(html.escape(a) for a in avisos) + '</div>')
        T.append('\nAVISOS\n' + '\n'.join('- ' + a for a in avisos))
    H.append(f'<div style="font-size:11px;color:{MUTED};padding:12px 0 18px">A planilha completa (painel, gráficos, equipe, semanas, resgate, agenda e caixa) fica na conversa do Claude desta execução.</div></div>')
    T.append('\nA planilha completa fica na conversa do Claude desta execução.')
    return {'assunto': assunto, 'html': ''.join(H), 'texto': '\n'.join(T)}

def page(em, gerado):
    """Página do painel (artifact) com o último relatório."""
    return f'''<title>Painel Encanto</title>
<meta name="claude:viewport-fit-sides" content="cover">
<style>
/* Página de leitura única no tema claro, igual ao e-mail: faixa escura no topo, relatório centralizado */
:root {{ --bg:#ECEAEF; --card:#FFFFFF; --ink:#2A2A33; --muted:#8A8A99; --accent:#D9644A; }}
body {{ background:var(--bg); color:var(--ink); font-family:Arial,Helvetica,sans-serif; padding-block:24px;
        padding-left:max(16px, env(safe-area-inset-left, 0px)); padding-right:max(16px, env(safe-area-inset-right, 0px)); }}
.wrap {{ max-width:700px; margin:0 auto; display:flex; flex-direction:column; gap:12px; }}
.meta {{ font-size:12px; color:var(--muted); }}
.card {{ background:var(--card); padding:12px; overflow-x:auto; }}
.card table {{ min-width:0; }}
h1 {{ font-size:13px; letter-spacing:.06em; text-transform:uppercase; color:var(--accent); margin:0; }}
</style>
<div class="wrap">
  <h1>Relatório Trinks · atualizado em {html.escape(gerado)}</h1>
  <div class="meta">{html.escape(em["assunto"])}. Atualiza de segunda a sábado por volta das 08h e das 14h.</div>
  <div class="card">{em["html"]}</div>
</div>'''
